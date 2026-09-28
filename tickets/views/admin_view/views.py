# tickets/views/admin_view/views.py
# ============================================================
# ADMIN_VIEW — Oversight (Read-Only + Limited Write)
# ============================================================

import csv
from datetime import timedelta

from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.http import HttpResponse, HttpResponseBadRequest
from django.utils import timezone
from django.db.models import Count, Q
from django.core.paginator import Paginator

from tickets.models import Ticket, TicketReply, Unit, Department

from .decorators import auditor_required, auditor_can


# ============================================================
# Helpers
# ============================================================
def _priority_value(ticket):
    p = getattr(ticket, 'priority', None)
    if isinstance(p, dict):
        return p.get('value') or ''
    return p or ''


def _priority_class(ticket):
    p = getattr(ticket, 'priority', None)
    if isinstance(p, dict):
        return p.get('class') or f"badge-priority-{_priority_value(ticket).lower()}"
    return f"badge-priority-{_priority_value(ticket).lower()}"


def _priority_display(ticket):
    v = _priority_value(ticket)
    mapping = {'Low': 'Low', 'Medium': 'Medium', 'High': 'High', 'Critical': 'Critical'}
    return mapping.get(v, v or '—')


def _status_value(ticket):
    s = getattr(ticket, 'status', None)
    if isinstance(s, dict):
        return s.get('value') or ''
    return s or ''


def _status_class(ticket):
    s = getattr(ticket, 'status', None)
    if isinstance(s, dict):
        return s.get('class') or f"badge-status-{_status_value(ticket).lower()}"
    return f"badge-status-{_status_value(ticket).lower()}"


def _status_display(ticket):
    v = _status_value(ticket)
    mapping = {
        'Open': 'Open',
        'Assigned': 'Assigned',
        'Hold': 'On Hold',
        'Escalated': 'Escalated',
        'Closed': 'Closed',
    }
    return mapping.get(v, v or '—')


# ============================================================
# Presets
# ============================================================
PRESETS = {
    'all':              [],
    'open':             ['Open'],
    'assigned':         ['Assigned'],
    'hold':             ['Hold'],
    'escalated':        ['Escalated'],
    'closed':           ['Closed'],
    'escalated_closed': ['Escalated', 'Closed'],
}


# ============================================================
# Filters
# ============================================================
def _apply_filters(qs, request):
    search         = request.GET.get('q', '').strip()
    ticket_number  = request.GET.get('ticket_number', '').strip()
    preset         = request.GET.get('preset', '').strip()

    selected_statuses    = request.GET.getlist('status')
    selected_priorities  = request.GET.getlist('priority')
    selected_units       = request.GET.getlist('unit')
    selected_departments = request.GET.getlist('department')
    selected_assignees   = request.GET.getlist('assignee')

    if preset and preset in PRESETS and PRESETS[preset]:
        if not selected_statuses:
            selected_statuses = PRESETS[preset]

    if selected_statuses:
        qs = qs.filter(status__in=selected_statuses)

    if selected_priorities:
        qs = qs.filter(priority__in=selected_priorities)

    if selected_units:
        qs = qs.filter(unit_id__in=selected_units)

    if selected_departments:
        qs = qs.filter(department_id__in=selected_departments)

    if selected_assignees:
        has_unassigned = '__unassigned__' in selected_assignees
        real_names = [a for a in selected_assignees if a != '__unassigned__']

        if has_unassigned and real_names:
            qs = qs.filter(
                Q(assigned_person__in=real_names)
                | Q(assigned_person__isnull=True)
                | Q(assigned_person='')
            )
        elif has_unassigned:
            qs = qs.filter(Q(assigned_person__isnull=True) | Q(assigned_person=''))
        else:
            qs = qs.filter(assigned_person__in=real_names)

    if ticket_number:
        qs = qs.filter(ticket_number__icontains=ticket_number)

    if search:
        qs = qs.filter(
            Q(subject__icontains=search)
            | Q(description__icontains=search)
            | Q(ticket_number__icontains=search)
        )

    return qs


# ============================================================
# DASHBOARD
# ============================================================
@auditor_required
def dashboard(request):
    tickets = Ticket.objects.all()
    now = timezone.now()
    week_ago = now - timedelta(days=7)

    # ---- KPI counts ----
    total           = tickets.count()
    open_count      = tickets.filter(status='Open').count()
    critical_count  = tickets.filter(priority='Critical').count()
    escalated_count = tickets.filter(status='Escalated').count()
    overdue_count   = tickets.filter(
        target_date__lt=now,
        status__in=['Open', 'Assigned', 'Hold']
    ).count()
    closed_week     = tickets.filter(
        status='Closed',
        closed_at__gte=week_ago
    ).count()

    # ---- Trend: tickets per day, last 180 days ----
    days  = 180
    start = now - timedelta(days=days - 1)

    raw = (
        tickets
        .filter(created_at__gte=start)
        .values('created_at__date')
        .annotate(count=Count('id'))
    )
    by_day = {row['created_at__date']: row['count'] for row in raw}

    trend_labels = []
    trend_values = []
    for i in range(days - 1, -1, -1):
        day = (now - timedelta(days=i)).date()
        trend_labels.append(day.strftime('%b %d'))
        trend_values.append(by_day.get(day, 0))

    # ---- Data for the chart (JSON-serializable) ----
    av_dashboard_data = {
        'trend': {
            'labels': trend_labels,
            'values': trend_values,
        },
    }

    context = {
        'page_title': 'Oversight Dashboard',
        'total': total,
        'open_count': open_count,
        'critical_count': critical_count,
        'escalated_count': escalated_count,
        'overdue_count': overdue_count,
        'closed_week': closed_week,
        'recent_tickets': tickets.order_by('-created_at')[:10],
        'av_dashboard_data': av_dashboard_data,
    }
    return render(request, 'admin_view/dashboard.html', context)


# ============================================================
# TICKETS PAGE
# ============================================================
@auditor_required
def tickets_page(request, ticket_id=None):
    base_qs = Ticket.objects.select_related(
        'unit', 'department', 'created_by_user'
    ).all()

    qs = _apply_filters(base_qs, request)

    sort = request.GET.get('sort', '-created_at')
    allowed_sorts = [
        '-created_at', 'created_at',
        'status', '-status',
        'priority', '-priority',
        'target_date', '-target_date',
    ]
    qs = qs.order_by(sort if sort in allowed_sorts else '-created_at')

    paginator = Paginator(qs, 10)
    page = paginator.get_page(request.GET.get('page'))

    selected_statuses    = request.GET.getlist('status')
    selected_priorities  = request.GET.getlist('priority')
    selected_units       = request.GET.getlist('unit')
    selected_departments = request.GET.getlist('department')
    selected_assignees   = request.GET.getlist('assignee')
    preset               = request.GET.get('preset', '')

    all_units = Unit.objects.filter(is_active=True).order_by('code')

    if selected_units:
        all_departments = Department.objects.filter(
            unit_id__in=selected_units
        ).order_by('name')
        disabled_departments = set(
            Department.objects.exclude(
                unit_id__in=selected_units
            ).values_list('id', flat=True)
        )
    else:
        all_departments = Department.objects.none()
        disabled_departments = set()

    if selected_units:
        assignee_qs = (
            Ticket.objects
            .exclude(assigned_person__isnull=True)
            .exclude(assigned_person='')
            .filter(unit_id__in=selected_units)
        )
        if selected_departments:
            assignee_qs = assignee_qs.filter(department_id__in=selected_departments)
        all_assignees = (
            assignee_qs.values_list('assigned_person', flat=True)
            .distinct()
            .order_by('assigned_person')
        )
    else:
        all_assignees = []

    status_counts = {
        s: base_qs.filter(status=s).count()
        for s in ['Open', 'Assigned', 'Hold', 'Escalated', 'Closed']
    }
    priority_counts = {
        p: base_qs.filter(priority=p).count()
        for p in ['Critical', 'High', 'Medium', 'Low']
    }

    # ---- Permission flags ----
    can_export_reports = request.user.has_perm('tickets.export_ticket_reports')

    # ---- Selected ticket (detail view) ----
    selected_ticket = None
    replies = []
    can_comment = can_change_priority = False
    priority_value = priority_class = priority_display = None
    status_value = status_class = status_display = None
    history = []

    if ticket_id:
        selected_ticket = get_object_or_404(
            Ticket.objects.select_related('unit', 'department', 'created_by_user'),
            pk=ticket_id
        )
        replies = selected_ticket.replies.select_related('author').order_by('created_at')

        user = request.user
        can_comment         = user.has_perm('tickets.add_comment')
        can_change_priority = user.has_perm('tickets.change_ticket_priority')

        priority_value   = _priority_value(selected_ticket)
        priority_class   = _priority_class(selected_ticket)
        priority_display = _priority_display(selected_ticket)

        status_value   = _status_value(selected_ticket)
        status_class   = _status_class(selected_ticket)
        status_display = _status_display(selected_ticket)

        # ---------- AUDIT HISTORY ----------
        log_qs = None
        for related_name in ('logs', 'history', 'audit_logs', 'ticketlog_set'):
            try:
                log_qs = getattr(selected_ticket, related_name).all()
                break
            except AttributeError:
                log_qs = None

        if log_qs is not None:
            try:
                log_qs = log_qs.order_by('-timestamp')
            except Exception:
                try:
                    log_qs = log_qs.order_by('-created_at')
                except Exception:
                    pass

            for entry in log_qs:
                action = (
                    getattr(entry, 'action', None)
                    or getattr(entry, 'event', None)
                    or getattr(entry, 'action_type', None)
                    or 'Event'
                )
                performed_by = (
                    getattr(entry, 'performed_by', None)
                    or getattr(entry, 'user', None)
                    or getattr(entry, 'actor', None)
                )
                remarks = (
                    getattr(entry, 'remarks', None)
                    or getattr(entry, 'message', None)
                    or getattr(entry, 'note', None)
                    or ''
                )
                ts = (
                    getattr(entry, 'timestamp', None)
                    or getattr(entry, 'created_at', None)
                )
                history.append({
                    'action': str(action),
                    'performed_by': str(performed_by) if performed_by else 'System',
                    'remarks': str(remarks) if remarks else '',
                    'timestamp': ts,
                })

        if not history:
            history.append({
                'action': 'Created',
                'performed_by': (
                    selected_ticket.created_by_user.username
                    if selected_ticket.created_by_user else 'System'
                ),
                'remarks': f'Ticket #{selected_ticket.ticket_number} created.',
                'timestamp': selected_ticket.created_at,
            })
            if getattr(selected_ticket, 'escalated_at', None):
                history.append({
                    'action': 'Escalated',
                    'performed_by': 'System',
                    'remarks': getattr(selected_ticket, 'vendor_ticket_number', '') or '',
                    'timestamp': selected_ticket.escalated_at,
                })
            if getattr(selected_ticket, 'closed_at', None):
                history.append({
                    'action': 'Closed',
                    'performed_by': getattr(selected_ticket, 'closed_by', '') or 'System',
                    'remarks': getattr(selected_ticket, 'closing_remarks', '') or '',
                    'timestamp': selected_ticket.closed_at,
                })

    context = {
        'page_title': 'All Tickets (Oversight)',
        'can_export_reports': can_export_reports,
        'page_obj': page,
        'total_count': qs.count(),
        'all_units': all_units,
        'all_departments': all_departments,
        'all_assignees': all_assignees,
        'status_counts': status_counts,
        'priority_counts': priority_counts,
        'disabled_departments': disabled_departments,
        'selected_statuses': selected_statuses,
        'selected_priorities': selected_priorities,
        'selected_units': selected_units,
        'selected_departments': selected_departments,
        'selected_assignees': selected_assignees,
        'has_units_selected': bool(selected_units),
        'has_departments_selected': bool(selected_departments),
        'current_filters': {
            'q': request.GET.get('q', ''),
            'ticket_number': request.GET.get('ticket_number', ''),
            'preset': preset,
            'sort': sort,
        },
        'selected_ticket': selected_ticket,
        'replies': replies,
        'history': history,
        'can_comment': can_comment,
        'can_change_priority': can_change_priority,
        'priority_value': priority_value,
        'priority_class': priority_class,
        'priority_display': priority_display,
        'status_value': status_value,
        'status_class': status_class,
        'status_display': status_display,
    }

    if ticket_id:
        return render(request, 'admin_view/ticket_detail.html', context)
    return render(request, 'admin_view/ticket_list.html', context)


# ============================================================
# ACTIONS
# ============================================================
@auditor_can('tickets.add_comment')
def add_comment(request, ticket_id):
    ticket = get_object_or_404(Ticket, pk=ticket_id)

    if request.method != 'POST':
        return redirect('admin_view:ticket_detail', ticket_id=ticket.id)

    body = request.POST.get('body', '').strip()
    if not body:
        messages.error(request, 'Comment cannot be empty.')
        return redirect('admin_view:ticket_detail', ticket_id=ticket.id)

    TicketReply.objects.create(
        ticket=ticket,
        author=request.user,
        author_name=request.user.get_full_name() or request.user.username,
        author_role='Auditor',
        body=body,
    )

    messages.success(request, 'Comment posted.')
    return redirect('admin_view:ticket_detail', ticket_id=ticket.id)


@auditor_can('tickets.change_ticket_priority')
def change_priority(request, ticket_id):
    ticket = get_object_or_404(Ticket, pk=ticket_id)

    if request.method != 'POST':
        return redirect('admin_view:ticket_detail', ticket_id=ticket.id)

    new_priority = request.POST.get('priority')
    valid_priorities = ['Low', 'Medium', 'High', 'Critical']
    if new_priority not in valid_priorities:
        return HttpResponseBadRequest('Invalid priority')

    old_priority = _priority_value(ticket)
    ticket.priority = new_priority
    ticket.save(update_fields=['priority'])

    messages.success(
        request,
        f'Priority changed from {old_priority} to {new_priority}.'
    )
    return redirect('admin_view:ticket_detail', ticket_id=ticket.id)


# ============================================================
# SINGLE-TICKET EXPORTS
# ============================================================
def _single_ticket_rows(ticket):
    def _sv(v):
        if isinstance(v, dict):
            return v.get('value', '') or ''
        return v or ''

    return [
        ('Ticket #',      ticket.ticket_number),
        ('Subject',       ticket.subject),
        ('Employee',      getattr(ticket, 'employee_name', '') or ''),
        ('Employee ID',   getattr(ticket, 'employee_id', '') or ''),
        ('Email',         getattr(ticket, 'email', '') or ''),
        ('Mobile',        getattr(ticket, 'mobile', '') or ''),
        ('Unit',          ticket.unit.full_name if getattr(ticket, 'unit', None) else ''),
        ('Department',    ticket.department.name if getattr(ticket, 'department', None) else ''),
        ('Screen',        getattr(ticket, 'screen_number', '') or ''),
        ('Error Type',    getattr(ticket, 'error_type', '') or ''),
        ('Priority',      _sv(ticket.priority)),
        ('Status',        _sv(ticket.status)),
        ('Assigned To',   getattr(ticket, 'assigned_person', '') or '—'),
        ('Target Date',   ticket.target_date.strftime('%Y-%m-%d') if getattr(ticket, 'target_date', None) else '—'),
        ('Created At',    ticket.created_at.strftime('%Y-%m-%d %H:%M') if getattr(ticket, 'created_at', None) else ''),
        ('Last Updated',  ticket.updated_at.strftime('%Y-%m-%d %H:%M') if getattr(ticket, 'updated_at', None) else ''),
        ('Vendor Ticket', getattr(ticket, 'vendor_ticket_number', '') or '—'),
        ('Escalated At',  ticket.escalated_at.strftime('%Y-%m-%d %H:%M') if getattr(ticket, 'escalated_at', None) else '—'),
    ]


@auditor_required
def download_ticket_excel(request, pk):
    ticket = get_object_or_404(
        Ticket.objects.select_related('unit', 'department', 'created_by_user'),
        pk=pk
    )

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return HttpResponseBadRequest('openpyxl is not installed.')

    wb = Workbook()
    ws = wb.active
    ws.title = f'Ticket-{ticket.ticket_number}'

    def _sv(v):
        if isinstance(v, dict):
            return v.get('value', '') or ''
        return v or ''

    ws.merge_cells('A1:B1')
    ws['A1'] = f'Ticket #{ticket.ticket_number} — {ticket.subject}'
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='FF6B00', end_color='FF6B00', fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 26

    for label, value in _single_ticket_rows(ticket):
        ws.append([label, str(value)])

    ws.append([])
    ws.append(['Description', ''])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(['', getattr(ticket, 'description', '') or 'No description.'])

    if _sv(ticket.status) == 'Closed':
        ws.append([])
        ws.append(['Closing Details', ''])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append(['Main Error Type', getattr(ticket, 'main_error_type', '') or '—'])
        ws.append(['Sub Error Type',  getattr(ticket, 'sub_error_type', '') or '—'])
        ws.append(['Closing Remarks', getattr(ticket, 'closing_remarks', '') or '—'])
        ws.append(['Closed By',       getattr(ticket, 'closed_by', '') or '—'])
        ws.append([
            'Closed At',
            ticket.closed_at.strftime('%Y-%m-%d %H:%M') if getattr(ticket, 'closed_at', None) else '—'
        ])

    replies = ticket.replies.all().order_by('created_at')
    if replies.exists():
        ws.append([])
        ws.append(['Conversation', ''])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for r in replies:
            when = r.created_at.strftime('%Y-%m-%d %H:%M') if getattr(r, 'created_at', None) else ''
            who = (
                getattr(r, 'author_name', '')
                or (r.author.username if getattr(r, 'author', None) else 'Unknown')
            )
            body = getattr(r, 'body', '') or getattr(r, 'message', '') or ''
            ws.append([f'{who} · {when}', body])

    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 70

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = (
        f'attachment; filename="ticket-{ticket.ticket_number}.xlsx"'
    )
    wb.save(response)
    return response


@auditor_required
def download_ticket_csv(request, pk):
    ticket = get_object_or_404(
        Ticket.objects.select_related('unit', 'department', 'created_by_user'),
        pk=pk
    )

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = (
        f'attachment; filename="ticket-{ticket.ticket_number}.csv"'
    )

    writer = csv.writer(response)
    writer.writerow(['Field', 'Value'])

    for label, value in _single_ticket_rows(ticket):
        writer.writerow([label, str(value)])

    writer.writerow([])
    writer.writerow(['Description', ''])
    writer.writerow(['', getattr(ticket, 'description', '') or 'No description.'])

    def _sv(v):
        if isinstance(v, dict):
            return v.get('value', '') or ''
        return v or ''

    if _sv(ticket.status) == 'Closed':
        writer.writerow([])
        writer.writerow(['Closing Details', ''])
        writer.writerow(['Main Error Type', getattr(ticket, 'main_error_type', '') or '—'])
        writer.writerow(['Sub Error Type',  getattr(ticket, 'sub_error_type', '') or '—'])
        writer.writerow(['Closing Remarks', getattr(ticket, 'closing_remarks', '') or '—'])
        writer.writerow(['Closed By',       getattr(ticket, 'closed_by', '') or '—'])
        writer.writerow([
            'Closed At',
            ticket.closed_at.strftime('%Y-%m-%d %H:%M') if getattr(ticket, 'closed_at', None) else '—'
        ])

    replies = ticket.replies.all().order_by('created_at')
    if replies.exists():
        writer.writerow([])
        writer.writerow(['Conversation', ''])
        for r in replies:
            when = r.created_at.strftime('%Y-%m-%d %H:%M') if getattr(r, 'created_at', None) else ''
            who = (
                getattr(r, 'author_name', '')
                or (r.author.username if getattr(r, 'author', None) else 'Unknown')
            )
            body = getattr(r, 'body', '') or getattr(r, 'message', '') or ''
            writer.writerow([f'{who} · {when}', body])

    return response


# ============================================================
# REPORTS
# ============================================================
@auditor_required
def reports(request):
    tickets = Ticket.objects.all()

    by_status = (
        tickets.values('status')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    by_priority = (
        tickets.values('priority')
        .annotate(count=Count('id'))
        .order_by('-count')
    )
    by_unit = (
        tickets.values('unit__code')
        .annotate(count=Count('id'))
        .order_by('-count')[:10]
    )

    context = {
        'page_title': 'Oversight Reports',
        'by_status': list(by_status),
        'by_priority': list(by_priority),
        'by_unit': list(by_unit),
        'can_export_reports': request.user.has_perm('tickets.export_ticket_reports'),
    }
    return render(request, 'admin_view/reports.html', context)


# ============================================================
# LIST EXPORTS
# ============================================================
@auditor_required
def export_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="gplast_tickets.csv"'

    writer = csv.writer(response)
    writer.writerow([
        'Ticket #', 'Subject', 'Status', 'Priority',
        'Unit', 'Department', 'Assigned Person',
        'Created By', 'Created', 'Target Date',
        'Description'
    ])

    tickets = Ticket.objects.select_related('unit', 'department', 'created_by_user').all()
    tickets = _apply_filters(tickets, request).order_by('-created_at')

    for t in tickets:
        writer.writerow([
            t.ticket_number,
            t.subject,
            _status_value(t),
            _priority_value(t),
            t.unit.full_name if t.unit else '',
            t.department.name if t.department else '',
            t.assigned_person or '',
            t.created_by_user.username if t.created_by_user else '',
            t.created_at.strftime('%Y-%m-%d %H:%M') if t.created_at else '',
            t.target_date.strftime('%Y-%m-%d') if t.target_date else '',
            (t.description or '').replace('\r\n', ' ').replace('\n', ' '),
        ])

    return response


@auditor_required
def export_xlsx(request):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return HttpResponseBadRequest('openpyxl is not installed.')

    wb = Workbook()
    ws = wb.active
    ws.title = 'Tickets'

    headers = [
        'Ticket #', 'Subject', 'Status', 'Priority',
        'Unit', 'Department', 'Assigned Person',
        'Created By', 'Created', 'Target Date',
        'Description'
    ]
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='FF6B00', end_color='FF6B00', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    tickets = Ticket.objects.select_related('unit', 'department', 'created_by_user').all()
    tickets = _apply_filters(tickets, request).order_by('-created_at')

    for t in tickets:
        ws.append([
            t.ticket_number,
            t.subject,
            _status_value(t),
            _priority_value(t),
            t.unit.full_name if t.unit else '',
            t.department.name if t.department else '',
            t.assigned_person or '',
            t.created_by_user.username if t.created_by_user else '',
            t.created_at.strftime('%Y-%m-%d %H:%M') if t.created_at else '',
            t.target_date.strftime('%Y-%m-%d') if t.target_date else '',
            t.description or '',
        ])

    fixed_widths = {
        'A': 10, 'B': 40, 'C': 12, 'D': 12, 'E': 30,
        'F': 20, 'G': 20, 'H': 16, 'I': 18, 'J': 14, 'K': 60,
    }
    for col, w in fixed_widths.items():
        ws.column_dimensions[col].width = w

    ws.row_dimensions[1].height = 22

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="gplast_tickets.xlsx"'
    wb.save(response)
    return response